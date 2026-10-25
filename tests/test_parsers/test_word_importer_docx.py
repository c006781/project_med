# tests/test_parsers/test_word_importer_docx.py

"""
Импорт синтетического ``.docx`` во временную БД (``parsers/word_importer.py``).

Задача плана Ф1.1.6 (2026-10-25) и исправления прохода 1, часть 2: W2 (предупреждение
доходит до результатов, журнала и колбэка), W7 (отбор файлов в ``batch_parse`` из окна,
пустая папка — пара), W8 (временная папка фото — не в папке документов), W10 (текст
ошибки в колбэке), Ф1.1.7 (``pars_start_args``).

Документы создаются в тесте python-docx во временной папке; БД — в памяти (``tests/conftest.py``),
хранилище фото — ``tmp_path``. Без сети. Образцы настоящих документов — после уточнения у
владельца.
"""

import os
import struct
import zlib
from datetime import date

import pytest
from docx import Document

import parsers.word_importer as word_importer
from parsers.word_importer import (
    batch_parse,
    format_results,
    pars_start,
    pars_start_args,
    parse_one_file,
)


# Заголовки столбцов — как в шаблоне (образцы doc/)
TEMPLATE_HEADERS = [
    'дата приема',
    'причина обращения',
    'выполненная процедура',
    'Рекомендации и дата сл. приёма',
    'Стоимость процедуры',
    'Примечание',
    'Фото',
]


def _png_bytes() -> bytes:
    """
    Минимальный PNG 1×1 (белый пиксель) — рисунок для ячейки фото.

    Returns:
        bytes: Содержимое PNG-файла.
    """

    def chunk(kind: bytes, data: bytes) -> bytes:
        """Блок PNG: длина, тип, данные, CRC."""
        return (
            struct.pack('>I', len(data))
            + kind
            + data
            + struct.pack('>I', zlib.crc32(kind + data) & 0xFFFFFFFF)
        )

    header = struct.pack('>IIBBBBB', 1, 1, 8, 2, 0, 0, 0)     # 1×1, 8 бит, RGB
    pixels = zlib.compress(b'\x00\xff\xff\xff')               # фильтр 0 + белый пиксель
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', header) + chunk(b'IDAT', pixels) + chunk(b'IEND', b'')


def make_docx(
    path        : str,                 # куда сохранить .docx
    header_line : str,                 # шапка пациента (строка до таблицы)
    rows        : list,                # строки таблицы: списки из 6 текстов (без фото)
    photo_rows  : tuple = (),          # номера строк (с 0), в ячейку фото которых вставить рисунок
    table       : bool = True,         # False — документ без таблицы
) -> str:
    """
    Создаёт синтетический документ в формате шаблона: шапка, таблица 7 столбцов.

    Args:
        path (str): Путь к создаваемому файлу.
        header_line (str): Строка шапки пациента.
        rows (list): Строки приёмов — по 6 текстов (дата … примечание).
        photo_rows (tuple): Номера строк, в ячейку «Фото» которых вставить рисунок.
        table (bool): Создавать ли таблицу.

    Returns:
        str: Путь к файлу.

    Пример:
        >>> make_docx(str(tmp_path / 'а.docx'), 'Иванов Иван 1990', [['14.03.26', 'а', '', '', '', '']])
    """

    doc = Document()
    doc.add_paragraph(header_line)

    if table:
        word_table = doc.add_table(
            rows = 1 + len(rows),
            cols = len(TEMPLATE_HEADERS),
        )
        for idx, title in enumerate(TEMPLATE_HEADERS):
            word_table.rows[0].cells[idx].text = title

        png_path = os.path.join(
            os.path.dirname(path),
            '_pixel.png',
        )
        with open(png_path, 'wb') as f:
            f.write(_png_bytes())

        for row_idx, values in enumerate(rows, start=1):
            for col_idx, value in enumerate(values):
                word_table.rows[row_idx].cells[col_idx].text = value
            if row_idx - 1 in photo_rows:
                cell = word_table.rows[row_idx].cells[len(TEMPLATE_HEADERS) - 1]
                cell.paragraphs[0].add_run().add_picture(png_path)

        os.remove(png_path)

    doc.save(path)
    return path


ROWS = [
    ['14.03.26', 'Тест  1 1\n причина обращения', 'процедура 1', 'рекомендации 1', '1000', 'примечание 1'],
    ['14.04.2026', 'причина 2', 'процедура 2', '', '', ''],
]


@pytest.fixture
def services(database, patient_service, appointment_service, photo_service):
    """Сервисы и БД (в памяти) для импорта — как их собирает приложение."""
    return {
        'db'                  : database,
        'patient_service'     : patient_service,
        'appointment_service' : appointment_service,
        'photo_service'       : photo_service,
    }


def _patients(patient_service):
    """Все пациенты из БД (DTO)."""
    return patient_service.get_all()


# ----------------------------------------------------------------------
# Ф1.1.6 — импорт синтетического документа
# ----------------------------------------------------------------------

def test_import_synthetic_docx(tmp_path, db_session, services):
    """Пациент, приёмы, фото — из документа в БД; тексты очищены (W1), телефон приведён."""
    docs = tmp_path / 'docs'
    docs.mkdir()
    path = make_docx(
        str(docs / 'Иванов.docx'),
        'Иванов  Иван Иванович 1990 8 (3842) 12-34-56 тётя Анна',
        ROWS,
        photo_rows = (0,),
    )
    warnings = []

    assert parse_one_file(
        db_session,
        path,
        services['patient_service'],
        services['appointment_service'],
        services['photo_service'],
        warnings = warnings,
    ) is True

    [patient] = _patients(services['patient_service'])
    assert (patient.last_name, patient.first_name, patient.middle_name) == ('Иванов', 'Иван', 'Иванович')
    assert patient.birth_date == date(1990, 1, 1)
    assert patient.phone == '+7 (3842) 12-34-56'
    assert patient.description_text == 'тётя Анна'
    assert warnings == []

    appointments = sorted(
        services['appointment_service'].get_appointments_by_patient(patient.id),
        key = lambda a: a.date,
    )
    assert [a.date for a in appointments] == [date(2026, 3, 14), date(2026, 4, 14)]
    assert appointments[0].reason_text == 'Тест 1 1\nпричина обращения'
    assert appointments[0].cost_procedure_text == '1000'
    assert appointments[1].recommendations_text is None

    photos = services['photo_service'].get_photos_for_appointment(appointments[0].id)
    assert len(photos) == 1
    assert services['photo_service'].get_photos_for_appointment(appointments[1].id) == []

    # W8: в папке документов ничего лишнего не осталось
    assert sorted(os.listdir(docs)) == ['Иванов.docx']


# ----------------------------------------------------------------------
# W2 — предупреждение о телефоне доходит до результатов, журнала, колбэка
# ----------------------------------------------------------------------

def test_w2_batch_reports_unrecognized_phone(tmp_path, services):
    """W2: пациент импортирован без телефона, предупреждение — в результатах, журнале, колбэке."""
    make_docx(
        str(tmp_path / 'а.docx'),
        'Петров Пётр Петрович 1985 12-34-56',
        ROWS,
    )
    updates = []

    results, log_dir = batch_parse(
        str(tmp_path),
        progress_callback_update = lambda *args: updates.append(args),
        **services,
    )

    assert results['success'] == 1
    assert len(results['warnings']) == 1
    assert results['warnings'][0]['file_name'] == 'а.docx'
    assert '12-34-56' in results['warnings'][0]['warning']

    [patient] = _patients(services['patient_service'])
    assert patient.phone is None

    assert updates[-1][:2] == ('а.docx', 'success')
    assert '12-34-56' in updates[-1][2]

    with open(os.path.join(log_dir, 'parser.log'), encoding='utf-8') as f:
        [line] = f.read().splitlines()
    fields = line.split('\t')
    assert fields[4] == 'success'
    assert '12-34-56' in fields[5]


# ----------------------------------------------------------------------
# W7 — batch_parse без списка файлов (как из окна)
# ----------------------------------------------------------------------

def test_w7_window_call_skips_processed_and_temp(tmp_path, services):
    """W7: второй запуск без specific_files не импортирует тот же файл; «~$…» — пропуск."""
    make_docx(
        str(tmp_path / 'а.docx'),
        'Иванов Иван Иванович 1990',
        ROWS,
    )
    (tmp_path / '~$а.docx').write_bytes(b'')     # временный файл открытого Word

    first, _ = batch_parse(
        str(tmp_path),
        **services,
    )
    second, second_log_dir = batch_parse(
        str(tmp_path),
        **services,
    )

    assert (first['total'], first['success']) == (1, 1)
    assert second['total'] == 0
    assert second_log_dir is None
    assert len(_patients(services['patient_service'])) == 1


def test_w7_empty_folder_returns_pair(tmp_path, services):
    """W7: в папке нет .docx — пара (пустые результаты, None), как ждёт окно."""
    results, log_dir = batch_parse(
        str(tmp_path),
        **services,
    )

    assert results == {
        'total'         : 0,
        'success'       : 0,
        'failed'        : 0,
        'success_files' : [],
        'error_files'   : [],
        'warnings'      : [],
    }
    assert log_dir is None
    assert not (tmp_path / 'logs_parser').exists()


# ----------------------------------------------------------------------
# W8 — временная папка фото
# ----------------------------------------------------------------------

def test_w8_temp_dir_outside_documents(tmp_path, db_session, services, monkeypatch):
    """W8: папка пользователя temp_images цела; временная папка — своя и удалена."""
    user_dir = tmp_path / 'temp_images'
    user_dir.mkdir()
    (user_dir / 'моё.txt').write_text('данные пользователя', encoding='utf-8')
    path = make_docx(
        str(tmp_path / 'а.docx'),
        'Иванов Иван Иванович 1990',
        ROWS,
        photo_rows = (0, 1),
    )

    created = []
    real_mkdtemp = word_importer.tempfile.mkdtemp

    def recording_mkdtemp(*args, **kwargs):
        """Запоминает созданные временные папки."""
        created.append(real_mkdtemp(*args, **kwargs))
        return created[-1]

    monkeypatch.setattr(word_importer.tempfile, 'mkdtemp', recording_mkdtemp)

    parse_one_file(
        db_session,
        path,
        services['patient_service'],
        services['appointment_service'],
        services['photo_service'],
    )

    assert (user_dir / 'моё.txt').read_text(encoding='utf-8') == 'данные пользователя'
    assert len(created) == 1                       # одна папка на файл, не на ячейку
    assert not os.path.exists(created[0])          # и она удалена
    assert os.path.commonpath([created[0], str(tmp_path)]) != str(tmp_path)


# ----------------------------------------------------------------------
# W10 — текст ошибки в колбэке
# ----------------------------------------------------------------------

def test_w10_failed_callback_gets_text(tmp_path, services):
    """W10: у файла с ошибкой колбэк получает текст ошибки (строку), журнал — failed."""
    make_docx(
        str(tmp_path / 'без таблицы.docx'),
        'Иванов Иван Иванович 1990',
        [],
        table = False,
    )
    updates = []

    results, _ = batch_parse(
        str(tmp_path),
        progress_callback_update = lambda *args: updates.append(args),
        **services,
    )

    assert results['failed'] == 1
    assert updates[-1][:2] == ('без таблицы.docx', 'failed')
    assert isinstance(updates[-1][2], str)
    assert 'таблица' in updates[-1][2]
    assert results['warnings'] == []


# ----------------------------------------------------------------------
# pars_start, pars_start_args (Ф1.1.7), format_results
# ----------------------------------------------------------------------

def test_pars_start_log_path_is_file(tmp_path, services, monkeypatch):
    """pars_start: log_path — путь к parser.log (раньше — папка журнала)."""
    make_docx(
        str(tmp_path / 'а.docx'),
        'Иванов Иван Иванович 1990',
        ROWS,
    )
    real_batch_parse = word_importer.batch_parse
    monkeypatch.setattr(
        word_importer,
        'batch_parse',
        lambda *args, **kwargs: real_batch_parse(*args, **kwargs, **services),
    )

    results = pars_start(str(tmp_path))

    assert results['success'] == 1
    assert results['log_path'] == os.path.join(str(tmp_path), 'logs_parser', 'parser.log')
    assert os.path.isfile(results['log_path'])


@pytest.fixture
def captured_pars_start(monkeypatch):
    """Подмена pars_start: запоминает параметры вызова, в БД не пишет."""
    calls = []

    def fake_pars_start(**kwargs):
        """Запоминает параметры, возвращает пустые результаты."""
        calls.append(kwargs)
        return {}

    monkeypatch.setattr(word_importer, 'pars_start', fake_pars_start)
    return calls


def test_pars_start_args_folder_from_command_line(tmp_path, captured_pars_start):
    """Ф1.1.7: --folder главнее папки из вызова; --files — список без пустых."""
    pars_start_args(
        folder = '/нет/такой/папки',
        argv   = ['--folder', str(tmp_path), '--files', 'а.docx, б.docx,', '--update_existing_patient'],
    )

    assert captured_pars_start == [
        {
            'folder'                  : str(tmp_path),
            'files'                   : ['а.docx', 'б.docx'],
            'update_existing_patient' : True,
        }
    ]


def test_pars_start_args_folder_from_call(tmp_path, captured_pars_start):
    """Без --folder — папка из вызова; без --files — None (все необработанные)."""
    pars_start_args(
        folder = str(tmp_path),
        argv   = [],
    )

    assert captured_pars_start[0]['folder'] == str(tmp_path)
    assert captured_pars_start[0]['files'] is None


@pytest.mark.parametrize(
    'argv',
    [
        [],                                   # папки нет нигде
        ['--folder', '/нет/такой/папки'],     # папки не существует
        ['--folder'],                         # --folder без значения (раньше IndexError)
    ],
)
def test_pars_start_args_bad_folder_exits(argv, captured_pars_start, capsys):
    """Нет папки — сообщение и выход с кодом 2, импорт не запускается."""
    with pytest.raises(SystemExit) as exc:
        pars_start_args(argv=argv)

    assert exc.value.code == 2
    assert captured_pars_start == []


def test_format_results():
    """Итоги: счётчики, ошибки, предупреждения, журнал."""
    text = format_results(
        {
            'total'       : 2,
            'success'     : 1,
            'failed'      : 1,
            'error_files' : [{'file_name': 'б.docx', 'error': ValueError('нет таблицы')}],
            'warnings'    : [{'file_name': 'а.docx', 'warning': 'Телефон не распознан'}],
            'log_path'    : 'doc/logs_parser/parser.log',
        }
    )

    assert text.splitlines() == [
        'Обработано файлов: 2',
        'Успешно: 1',
        'Ошибок: 1',
        '  б.docx: нет таблицы',
        'Предупреждения:',
        '  а.docx: Телефон не распознан',
        'Журнал: doc/logs_parser/parser.log',
    ]
